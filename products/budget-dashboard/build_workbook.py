"""Build the Budget Dashboard workbook sold as a digital download.

Run: python build_workbook.py [--blank]
  default  -> dist/Budget-Dashboard.xlsx (with sample data, for listing screenshots)
  --blank  -> dist/Budget-Dashboard-Blank.xlsx (what a buyer starts from)

Every number a buyer sees is a formula over the Settings and Transactions tabs,
so the file keeps working in Excel, Google Sheets, Numbers and LibreOffice.
Only Excel-2007-era functions are used for that reason.
"""

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BLANK = "--blank" in sys.argv
OUT = Path(__file__).parent / "dist" / (
    "Budget-Dashboard-Blank.xlsx" if BLANK else "Budget-Dashboard.xlsx"
)

YEAR = 2026
CAT_FIRST, CAT_LAST = 8, 47  # 40 category slots on Settings
TX_FIRST, TX_LAST = 6, 2005  # 2,000 transaction rows
GOAL_FIRST, GOAL_LAST = 6, 17
DEBT_FIRST, DEBT_LAST = 6, 17
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

# Palette: calm forest green with warm accents; reads well printed or on screen.
INK = "1F2A24"
MUTED = "6B7A71"
BRAND = "2F6B4F"
BRAND_DARK = "1E4A36"
BRAND_SOFT = "E6F0EA"
INPUT = "FFF7D6"
GOOD = "2E7D32"
BAD = "C62828"
LINE = "D5DED8"

FONT = "Arial"
CUR = '$#,##0.00;($#,##0.00);"-"'
CUR0 = '$#,##0;($#,##0);"-"'
PCT = '0%;-0%;"-"'

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def base(ws, widths, tab_color):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def title(ws, text, subtitle, span):
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 30
    ws["B2"] = text
    ws["B2"].font = f(20, True, BRAND_DARK)
    ws["B3"] = subtitle
    ws["B3"].font = f(10, color=MUTED, italic=True)
    for c in range(2, 2 + span):
        ws.cell(row=3, column=c).border = Border(bottom=Side(style="medium", color=BRAND))


def header(ws, row, labels, start_col=2):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 28


def body(cell, fmt=None, is_input=False, align=None):
    cell.font = f()
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if is_input:
        cell.fill = fill(INPUT)
    if align:
        cell.alignment = Alignment(horizontal=align, vertical="center")


wb = Workbook()

# ---------------------------------------------------------------- Start Here
ws = wb.active
ws.title = "Start Here"
base(ws, {"A": 3, "B": 4, "C": 92}, BRAND_DARK)
title(ws, "Budget Dashboard", "Your money, one honest page at a time.", 2)
steps = [
    ("1", "Open Settings. Enter the year you're budgeting and your starting account balance."),
    ("2", "Still on Settings, edit the category list: name, Income or Expense, and a monthly budget. "
          "Up to 40 categories. Rename or delete the examples freely."),
    ("3", "Log every transaction on the Transactions tab: date, description, category (pick from the "
          "dropdown) and the amount as a positive number. Type and month fill themselves in."),
    ("4", "Open Dashboard and pick a month from the dropdown at the top. Budget vs. actual, what's left, "
          "and your savings rate update instantly."),
    ("5", "Year View shows all twelve months side by side. Goals tracks what you're saving for, "
          "and Debt Payoff tells you how long each balance will take at your current payment."),
]
r = 5
ws.cell(row=r, column=2, value="How to use it").font = f(13, True, BRAND_DARK)
r += 1
for n, text in steps:
    ws.cell(row=r, column=2, value=n).font = f(12, True, BRAND)
    ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
    c = ws.cell(row=r, column=3, value=text)
    c.font = f(11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32
    r += 1

r += 1
ws.cell(row=r, column=2, value="Legend").font = f(13, True, BRAND_DARK)
r += 1
ws.cell(row=r, column=2).fill = fill(INPUT)
ws.cell(row=r, column=2).border = BOX
ws.cell(row=r, column=3, value="Yellow cells are yours to type in.").font = f(11)
r += 1
ws.cell(row=r, column=2).fill = fill("FFFFFF")
ws.cell(row=r, column=2).border = BOX
ws.cell(row=r, column=3, value="White cells are formulas. Leave them alone and they'll keep doing the math.").font = f(11)
r += 2
ws.cell(row=r, column=2, value="Tips").font = f(13, True, BRAND_DARK)
tips = [
    "Refunds: enter them as a negative amount in the category you originally spent in.",
    "Moving money to savings? Keep the Expense category named exactly 'Savings': it counts against your plan, and the Dashboard adds it back into your savings rate.",
    "Google Sheets: File > Import > Upload, then 'Replace spreadsheet'. Everything works, dropdowns included.",
    "The file ships with sample data so you can see it working. Clear the yellow cells on Transactions, "
    "Goals and Debt Payoff when you're ready to start (a blank copy is included too).",
]
for t in tips:
    r += 1
    ws.cell(row=r, column=2, value="•").font = f(11, color=BRAND)
    ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
    c = ws.cell(row=r, column=3, value=t)
    c.font = f(11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
r += 2
ws.cell(row=r, column=3, value="For personal use. Please don't resell or redistribute this file.").font = f(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Settings
st = wb.create_sheet("Settings")
base(st, {"A": 3, "B": 30, "C": 14, "D": 18, "E": 3, "F": 44}, BRAND)
title(st, "Settings", "Set these once. Everything else reads from here.", 3)

st["B5"], st["C5"] = "Budget year", YEAR
st["B6"], st["C6"] = "Starting balance", 0 if BLANK else 4250
for rr, fmt in ((5, "0"), (6, CUR)):
    st.cell(row=rr, column=2).font = f(10, True)
    body(st.cell(row=rr, column=3), fmt, is_input=True, align="right")
st["C5"].comment = Comment("Transactions dated in other years are ignored by the Dashboard.", "Budget Dashboard")

header(st, CAT_FIRST - 1, ["Category", "Type", "Monthly budget"])
categories = [
    ("Salary", "Income", 5200), ("Side income", "Income", 400), ("Other income", "Income", 0),
    ("Rent / mortgage", "Expense", 1650), ("Utilities", "Expense", 180), ("Internet & phone", "Expense", 110),
    ("Groceries", "Expense", 520), ("Dining out", "Expense", 220), ("Transport", "Expense", 240),
    ("Insurance", "Expense", 190), ("Health", "Expense", 80), ("Subscriptions", "Expense", 45),
    ("Shopping", "Expense", 150), ("Entertainment", "Expense", 100), ("Personal care", "Expense", 60),
    ("Gifts & giving", "Expense", 75), ("Debt payments", "Expense", 350), ("Savings", "Expense", 600),
    ("Miscellaneous", "Expense", 75),
]
dv_type = DataValidation(type="list", formula1='"Income,Expense"', allow_blank=True)
st.add_data_validation(dv_type)
for i in range(CAT_FIRST, CAT_LAST + 1):
    name, kind, budget = categories[i - CAT_FIRST] if i - CAT_FIRST < len(categories) else (None, None, None)
    st.cell(row=i, column=2, value=name)
    st.cell(row=i, column=3, value=kind)
    st.cell(row=i, column=4, value=budget)
    body(st.cell(row=i, column=2), is_input=True)
    body(st.cell(row=i, column=3), is_input=True, align="center")
    body(st.cell(row=i, column=4), CUR, is_input=True)
dv_type.add(f"C{CAT_FIRST}:C{CAT_LAST}")
st.freeze_panes = f"B{CAT_FIRST}"

st["F7"] = "Planned for a month"
st["F7"].font = f(11, True, BRAND_DARK)
plan = [
    ("Planned income", f'=SUMIFS(D{CAT_FIRST}:D{CAT_LAST},C{CAT_FIRST}:C{CAT_LAST},"Income")'),
    ("Planned spending", f'=SUMIFS(D{CAT_FIRST}:D{CAT_LAST},C{CAT_FIRST}:C{CAT_LAST},"Expense")'),
    ("Left unassigned", "=F9-F11"),
]
rr = 8
for label, formula in plan:
    st.cell(row=rr, column=6, value=label).font = f(9, color=MUTED)
    c = st.cell(row=rr + 1, column=6, value=formula)
    c.font = f(14, True)
    c.number_format = CUR
    rr += 2
st.conditional_formatting.add("F13", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=14, bold=True, color=BAD)))
st["F14"] = "Aim for zero: every dollar of income has a job."
st["F14"].font = f(9, color=MUTED, italic=True)

CAT = f"Settings!$B${CAT_FIRST}:$B${CAT_LAST}"
CAT_TYPE = f"Settings!$C${CAT_FIRST}:$C${CAT_LAST}"
BUDGET_YEAR = "Settings!$C$5"

# ---------------------------------------------------------------- Transactions
tx = wb.create_sheet("Transactions")
base(tx, {"A": 3, "B": 13, "C": 34, "D": 22, "E": 13, "F": 11, "G": 9, "H": 8}, "C98A2B")
title(tx, "Transactions", "One row per transaction. Amounts are positive; refunds negative.", 7)
header(tx, TX_FIRST - 1, ["Date", "Description", "Category", "Amount", "Type", "Month", "Year"])

sample_tx = [] if BLANK else [
    (date(YEAR, 1, 1), "Paycheck", "Salary", 2600), (date(YEAR, 1, 2), "Rent", "Rent / mortgage", 1650),
    (date(YEAR, 1, 4), "Trader Joe's", "Groceries", 86.40), (date(YEAR, 1, 6), "Electric bill", "Utilities", 92.15),
    (date(YEAR, 1, 8), "Gas station", "Transport", 48.20), (date(YEAR, 1, 10), "Streaming bundle", "Subscriptions", 22.99),
    (date(YEAR, 1, 11), "Thai takeout", "Dining out", 41.80), (date(YEAR, 1, 12), "Car insurance", "Insurance", 190),
    (date(YEAR, 1, 14), "Farmers market", "Groceries", 34.00), (date(YEAR, 1, 15), "Paycheck", "Salary", 2600),
    (date(YEAR, 1, 15), "Transfer to savings", "Savings", 600), (date(YEAR, 1, 16), "Credit card payment", "Debt payments", 350),
    (date(YEAR, 1, 18), "Logo design gig", "Side income", 450), (date(YEAR, 1, 19), "Costco run", "Groceries", 212.60),
    (date(YEAR, 1, 21), "Phone bill", "Internet & phone", 65), (date(YEAR, 1, 21), "Internet", "Internet & phone", 45),
    (date(YEAR, 1, 23), "Birthday gift for Sam", "Gifts & giving", 60), (date(YEAR, 1, 24), "Brunch", "Dining out", 58.25),
    (date(YEAR, 1, 25), "Concert tickets", "Entertainment", 120), (date(YEAR, 1, 26), "Pharmacy", "Health", 24.50),
    (date(YEAR, 1, 27), "Running shoes", "Shopping", 129.99), (date(YEAR, 1, 28), "Groceries", "Groceries", 143.70),
    (date(YEAR, 1, 29), "Haircut", "Personal care", 45), (date(YEAR, 1, 30), "Train pass", "Transport", 127),
    (date(YEAR, 2, 1), "Paycheck", "Salary", 2600), (date(YEAR, 2, 1), "Rent", "Rent / mortgage", 1650),
    (date(YEAR, 2, 3), "Groceries", "Groceries", 118.35), (date(YEAR, 2, 5), "Electric bill", "Utilities", 104.60),
    (date(YEAR, 2, 7), "Pizza night", "Dining out", 36.00), (date(YEAR, 2, 9), "Gas station", "Transport", 51.75),
    (date(YEAR, 2, 12), "Car insurance", "Insurance", 190), (date(YEAR, 2, 14), "Valentine's dinner", "Dining out", 132.40),
    (date(YEAR, 2, 15), "Paycheck", "Salary", 2600), (date(YEAR, 2, 15), "Transfer to savings", "Savings", 600),
    (date(YEAR, 2, 16), "Credit card payment", "Debt payments", 350), (date(YEAR, 2, 18), "Groceries", "Groceries", 164.20),
    (date(YEAR, 2, 20), "Phone + internet", "Internet & phone", 110), (date(YEAR, 2, 22), "Movie", "Entertainment", 32),
    (date(YEAR, 2, 24), "Returned jacket", "Shopping", -59.99), (date(YEAR, 2, 26), "Groceries", "Groceries", 97.10),
    (date(YEAR, 3, 1), "Paycheck", "Salary", 2600), (date(YEAR, 3, 1), "Rent", "Rent / mortgage", 1650),
    (date(YEAR, 3, 4), "Groceries", "Groceries", 131.45), (date(YEAR, 3, 7), "Electric bill", "Utilities", 88.30),
    (date(YEAR, 3, 9), "Tax refund", "Other income", 740), (date(YEAR, 3, 12), "Car insurance", "Insurance", 190),
    (date(YEAR, 3, 15), "Paycheck", "Salary", 2600), (date(YEAR, 3, 15), "Transfer to savings", "Savings", 900),
    (date(YEAR, 3, 16), "Credit card payment", "Debt payments", 500), (date(YEAR, 3, 20), "Groceries", "Groceries", 176.90),
]
dv_cat = DataValidation(type="list", formula1=CAT, allow_blank=True, showErrorMessage=True,
                        errorTitle="Unknown category", error="Pick a category from the list, or add it on Settings first.")
tx.add_data_validation(dv_cat)
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, error="Enter a date, e.g. 3/14/2026.")
tx.add_data_validation(dv_date)
for i in range(TX_FIRST, TX_LAST + 1):
    row = sample_tx[i - TX_FIRST] if i - TX_FIRST < len(sample_tx) else (None, None, None, None)
    for col, val, fmt, align in ((2, row[0], "mm/dd/yyyy", "center"), (3, row[1], None, None),
                                 (4, row[2], None, None), (5, row[3], CUR, None)):
        c = tx.cell(row=i, column=col, value=val)
        body(c, fmt, is_input=True, align=align)
    tx.cell(row=i, column=6, value=f'=IF(D{i}="","",IFERROR(INDEX({CAT_TYPE},MATCH(D{i},{CAT},0)),"?"))')
    tx.cell(row=i, column=7, value=f'=IF(B{i}="","",MONTH(B{i}))')
    tx.cell(row=i, column=8, value=f'=IF(B{i}="","",YEAR(B{i}))')
    for col in (6, 7, 8):
        body(tx.cell(row=i, column=col), align="center")
        tx.cell(row=i, column=col).font = f(color=MUTED)
dv_cat.add(f"D{TX_FIRST}:D{TX_LAST}")
dv_date.add(f"B{TX_FIRST}:B{TX_LAST}")
tx.freeze_panes = f"B{TX_FIRST}"
tx.auto_filter.ref = f"B{TX_FIRST - 1}:H{TX_LAST}"

TXR = lambda col: f"Transactions!${col}${TX_FIRST}:${col}${TX_LAST}"  # noqa: E731
T_CAT, T_AMT, T_TYPE, T_MON, T_YR = TXR("D"), TXR("E"), TXR("F"), TXR("G"), TXR("H")

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 1)
base(db, {"A": 3, "B": 26, "C": 12, "D": 14, "E": 14, "F": 14, "G": 12, "H": 3, "I": 20, "J": 20}, GOOD)
title(db, "Dashboard", "Pick a month. The rest is automatic.", 6)

db["B5"] = "Month"
db["B5"].font = f(10, True)
db["C5"] = "January"
body(db["C5"], is_input=True, align="center")
db["C5"].font = f(11, True, BRAND_DARK)
db.merge_cells("C5:D5")
dv_month = DataValidation(type="list", formula1='"' + ",".join(MONTHS) + '"', allow_blank=False)
db.add_data_validation(dv_month)
dv_month.add("C5")
db["F5"] = "Year"
db["F5"].font = f(10, True)
db["F5"].alignment = Alignment(horizontal="right")
db["G5"] = f"={BUDGET_YEAR}"
db["G5"].font = f(11, True, BRAND_DARK)
db["G5"].alignment = Alignment(horizontal="center")
# Hidden helper: month number for the SUMIFS below.
db["J5"] = '=MATCH(C5,{"January","February","March","April","May","June","July","August","September","October","November","December"},0)'
db["J5"].font = f(color="FFFFFF")
MON = "$J$5"

def month_sum(kind):
    return f'=SUMIFS({T_AMT},{T_TYPE},"{kind}",{T_MON},{MON},{T_YR},{BUDGET_YEAR})'

kpis = [
    ("B", "Income", month_sum("Income"), CUR),
    ("D", "Spending", month_sum("Expense"), CUR),
    ("F", "Left over", "=B9-D9", CUR),
]
for col, label, formula, fmt in kpis:
    db[f"{col}8"] = label
    db[f"{col}8"].font = f(9, True, MUTED)
    db[f"{col}9"] = formula
    db[f"{col}9"].font = f(18, True, BRAND_DARK)
    db[f"{col}9"].number_format = fmt
    db[f"{col}9"].alignment = Alignment(horizontal="left")
    db[f"{col}10"].border = Border(bottom=Side(style="medium", color=BRAND))
db["I8"] = "Savings rate"
db["I8"].font = f(9, True, MUTED)
db["I9"] = '=IF(B9<=0,"-",(F9+SUMIFS(' + T_AMT + ',' + T_CAT + ',"Savings",' + T_MON + ',' + MON + ',' + T_YR + ',' + BUDGET_YEAR + '))/B9)'
db["I9"].font = f(18, True, BRAND_DARK)
db["I9"].number_format = PCT
db["I9"].alignment = Alignment(horizontal="left")
db["I9"].comment = Comment("Money left over plus anything logged to the 'Savings' category, as a share of income.", "Budget Dashboard")
db["J8"] = "Balance now"
db["J8"].font = f(9, True, MUTED)
db["J9"] = (f'=Settings!$C$6+SUMIFS({T_AMT},{T_TYPE},"Income",{T_YR},{BUDGET_YEAR},{T_MON},"<="&{MON})'
            f'-SUMIFS({T_AMT},{T_TYPE},"Expense",{T_YR},{BUDGET_YEAR},{T_MON},"<="&{MON})')
db["J9"].font = f(18, True, BRAND_DARK)
db["J9"].number_format = CUR
db["J9"].alignment = Alignment(horizontal="left")
db["J9"].comment = Comment("Starting balance plus everything earned, minus everything spent, through the end of the selected month.", "Budget Dashboard")
for col in ("I", "J"):
    db[f"{col}10"].border = Border(bottom=Side(style="medium", color=BRAND))
db.conditional_formatting.add("F9", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=18, bold=True, color=BAD)))

D_FIRST = 13
header(db, D_FIRST - 1, ["Category", "Type", "Budget", "Actual", "Remaining", "% used"])
D_LAST = D_FIRST + (CAT_LAST - CAT_FIRST)
for k in range(CAT_LAST - CAT_FIRST + 1):
    r, s = D_FIRST + k, CAT_FIRST + k
    db.cell(row=r, column=2, value=f'=IF(Settings!$B{s}="","",Settings!$B{s})')
    db.cell(row=r, column=3, value=f'=IF(B{r}="","",Settings!$C{s})')
    db.cell(row=r, column=4, value=f'=IF(B{r}="","",Settings!$D{s})')
    db.cell(row=r, column=5, value=f'=IF(B{r}="","",SUMIFS({T_AMT},{T_CAT},B{r},{T_MON},{MON},{T_YR},{BUDGET_YEAR}))')
    # Remaining flips sign for income so "positive" always means "good".
    db.cell(row=r, column=6, value=f'=IF(B{r}="","",IF(C{r}="Income",E{r}-D{r},D{r}-E{r}))')
    db.cell(row=r, column=7, value=f'=IF(OR(B{r}="",N(D{r})=0),"",E{r}/D{r})')
    for col, fmt, align in ((2, None, None), (3, None, "center"), (4, CUR, None), (5, CUR, None), (6, CUR, None), (7, PCT, "center")):
        body(db.cell(row=r, column=col), fmt, align=align)
    if k % 2:
        for col in range(2, 8):
            db.cell(row=r, column=col).fill = fill("F6F9F7")
db.conditional_formatting.add(
    f"F{D_FIRST}:F{D_LAST}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color=BAD, bold=True)))
db.conditional_formatting.add(
    f"G{D_FIRST}:G{D_LAST}",
    FormulaRule(formula=[f'AND($C{D_FIRST}="Expense",ISNUMBER(G{D_FIRST}),G{D_FIRST}>1)'], font=Font(name=FONT, color=BAD, bold=True),
                fill=fill("FDECEA")))
db.conditional_formatting.add(
    f"E{D_FIRST}:E{D_LAST}", DataBarRule(start_type="num", start_value=0, end_type="max", color="8FBCA5"))
# Collapse unused category rows so the table stays tidy.
db.conditional_formatting.add(
    f"B{D_FIRST}:G{D_LAST}", FormulaRule(formula=[f'$B{D_FIRST}=""'], fill=fill("FFFFFF"), border=Border()))

# Totals for the chart and a sanity footer.
tot = D_LAST + 1
db.cell(row=tot, column=2, value="Total spending").font = f(10, True)
db.cell(row=tot, column=4, value=f'=SUMIFS(D{D_FIRST}:D{D_LAST},C{D_FIRST}:C{D_LAST},"Expense")')
db.cell(row=tot, column=5, value=f'=SUMIFS(E{D_FIRST}:E{D_LAST},C{D_FIRST}:C{D_LAST},"Expense")')
db.cell(row=tot, column=6, value=f"=D{tot}-E{tot}")
for col in (4, 5, 6):
    c = db.cell(row=tot, column=col)
    c.font = f(10, True)
    c.number_format = CUR
    c.border = Border(top=Side(style="medium", color=BRAND))
db.cell(row=tot, column=2).border = Border(top=Side(style="medium", color=BRAND))
db.cell(row=tot, column=3).border = Border(top=Side(style="medium", color=BRAND))
db.cell(row=tot, column=7).border = Border(top=Side(style="medium", color=BRAND))

# Unrecognised categories would silently drop out of every total; say so loudly.
db["I12"] = "Needs attention"
db["I12"].font = f(10, True, BRAND_DARK)
db["I13"] = f'=COUNTIF({T_TYPE},"?")'
db["I13"].number_format = '0" rows with unknown category";;"All categories recognised"'
db["I13"].font = f(9, color=MUTED)
db.conditional_formatting.add("I13", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=9, bold=True, color=BAD)))

# Chart feeds: expense rows only, so a big paycheck doesn't flatten every spending bar.
# Kept visible (far right, muted): hidden cells drop out of charts in some spreadsheet apps.
db["R11"] = "Chart data - no need to edit"
db["R11"].font = f(8, color=MUTED, italic=True)
db["R12"], db["S12"] = "Budget", "Actual"
for k in range(-1, CAT_LAST - CAT_FIRST + 1):
    r = D_FIRST + k
    if k >= 0:
        db.cell(row=r, column=18, value=f'=IF(C{r}="Expense",D{r},0)')
        db.cell(row=r, column=19, value=f'=IF(C{r}="Expense",E{r},0)')
    for col in (18, 19):
        db.cell(row=r, column=col).font = f(8, color=MUTED)
        db.cell(row=r, column=col).number_format = CUR0

chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "Spending: budget vs. actual"
chart.y_axis.numFmt = "$#,##0"
chart.y_axis.majorGridlines = None
chart.x_axis.scaling.orientation = "maxMin"
chart.add_data(Reference(db, min_col=18, max_col=19, min_row=D_FIRST - 1, max_row=D_FIRST + len(categories) - 1), titles_from_data=True)
chart.set_categories(Reference(db, min_col=2, min_row=D_FIRST, max_row=D_FIRST + len(categories) - 1))
chart.series[0].graphicalProperties.solidFill = "C9DDD1"
chart.series[0].graphicalProperties.line.noFill = True
chart.series[1].graphicalProperties.solidFill = BRAND
chart.series[1].graphicalProperties.line.noFill = True
chart.height, chart.width = 15, 12
chart.legend.position = "b"
db.add_chart(chart, "I15")
db.freeze_panes = "A11"

# ---------------------------------------------------------------- Year View
yv = wb.create_sheet("Year View", 2)
widths = {"A": 3, "B": 24, "C": 11}
for j in range(12):
    widths[chr(ord("D") + j)] = 11
widths["P"] = 13
base(yv, widths, BRAND)
title(yv, "Year View", "Twelve months side by side.", 15)
header(yv, 5, ["Category", "Type"] + [m[:3] for m in MONTHS] + ["Year total"])
Y_FIRST = 6
for k in range(CAT_LAST - CAT_FIRST + 1):
    r, s = Y_FIRST + k, CAT_FIRST + k
    yv.cell(row=r, column=2, value=f'=IF(Settings!$B{s}="","",Settings!$B{s})')
    yv.cell(row=r, column=3, value=f'=IF(B{r}="","",Settings!$C{s})')
    body(yv.cell(row=r, column=2))
    body(yv.cell(row=r, column=3), align="center")
    for m in range(12):
        c = yv.cell(row=r, column=4 + m,
                    value=f'=IF($B{r}="","",SUMIFS({T_AMT},{T_CAT},$B{r},{T_MON},{m + 1},{T_YR},{BUDGET_YEAR}))')
        body(c, CUR0)
    body(yv.cell(row=r, column=16, value=f'=IF(B{r}="","",SUM(D{r}:O{r}))'), CUR0)
    yv.cell(row=r, column=16).font = f(bold=True)
Y_LAST = Y_FIRST + (CAT_LAST - CAT_FIRST)
yv.conditional_formatting.add(f"B{Y_FIRST}:P{Y_LAST}", FormulaRule(formula=[f'$B{Y_FIRST}=""'], fill=fill("FFFFFF"), border=Border()))

s_row = Y_LAST + 2
for offset, (label, formula) in enumerate((
    ("Income", 'SUMIFS({c}{f}:{c}{l},$C${f}:$C${l},"Income")'),
    ("Spending", 'SUMIFS({c}{f}:{c}{l},$C${f}:$C${l},"Expense")'),
    ("Left over", "{c}{inc}-{c}{exp}"),
)):
    r = s_row + offset
    c = yv.cell(row=r, column=2, value=label)
    c.font = f(10, True, "FFFFFF" if offset < 2 else BRAND_DARK)
    for j in range(13):
        col = chr(ord("D") + j)
        cell = yv.cell(row=r, column=4 + j,
                       value="=" + formula.format(c=col, f=Y_FIRST, l=Y_LAST, inc=s_row, exp=s_row + 1))
        cell.number_format = CUR0
        cell.font = f(10, True, "FFFFFF" if offset < 2 else BRAND_DARK)
    fill_color = BRAND if offset == 0 else ("B5543C" if offset == 1 else BRAND_SOFT)
    for col in range(2, 17):
        yv.cell(row=r, column=col).fill = fill(fill_color)
yv.conditional_formatting.add(f"D{s_row + 2}:P{s_row + 2}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color=BAD)))

ych = BarChart()
ych.style = 10
ych.title = "Income vs. spending by month"
ych.y_axis.numFmt = "$#,##0"
ych.y_axis.majorGridlines = None
ych.add_data(Reference(yv, min_col=3, max_col=15, min_row=s_row, max_row=s_row + 1), from_rows=True, titles_from_data=True)
ych.set_categories(Reference(yv, min_col=4, max_col=15, min_row=5))
# titles_from_data with from_rows takes column C as the title; C is blank, so name them explicitly.
from openpyxl.chart.series import SeriesLabel  # noqa: E402
ych.series[0].tx = SeriesLabel(v="Income")
ych.series[1].tx = SeriesLabel(v="Spending")
ych.series[0].graphicalProperties.solidFill = BRAND
ych.series[1].graphicalProperties.solidFill = "D98C6F"
ych.height, ych.width = 9, 26
ych.legend.position = "b"
yv.add_chart(ych, f"B{s_row + 5}")
yv.freeze_panes = "D6"

# ---------------------------------------------------------------- Goals
gl = wb.create_sheet("Goals")
base(gl, {"A": 3, "B": 28, "C": 14, "D": 14, "E": 12, "F": 14, "G": 14, "H": 16}, "7C5CBF")
title(gl, "Savings Goals", "What you're saving for, and what it takes each month to get there.", 7)
header(gl, GOAL_FIRST - 1, ["Goal", "Target", "Saved so far", "Progress", "Still needed", "Target date", "Save per month"])
sample_goals = [] if BLANK else [
    ("Emergency fund (3 months)", 9000, 4250, date(YEAR + 1, 6, 30)),
    ("Summer trip", 2400, 900, date(YEAR + 1, 7, 1)),
    ("New laptop", 1600, 300, date(YEAR + 1, 3, 1)),
]
for i in range(GOAL_FIRST, GOAL_LAST + 1):
    g = sample_goals[i - GOAL_FIRST] if i - GOAL_FIRST < len(sample_goals) else (None, None, None, None)
    for col, val, fmt in ((2, g[0], None), (3, g[1], CUR0), (4, g[2], CUR0), (7, g[3], "mm/dd/yyyy")):
        body(gl.cell(row=i, column=col, value=val), fmt, is_input=True)
    body(gl.cell(row=i, column=5, value=f'=IF(N(C{i})=0,"",MIN(1,D{i}/C{i}))'), PCT, align="center")
    body(gl.cell(row=i, column=6, value=f'=IF(B{i}="","",MAX(0,C{i}-D{i}))'), CUR0)
    # Whole months left, counting the current one; past-due goals ask for the full remainder now.
    body(gl.cell(row=i, column=8, value=(
        f'=IF(OR(B{i}="",G{i}=""),"",IF(F{i}=0,0,F{i}/MAX(1,(YEAR(G{i})-YEAR(TODAY()))*12+MONTH(G{i})-MONTH(TODAY())+1)))')), CUR0)
    gl.cell(row=i, column=8).font = f(bold=True)
gl.conditional_formatting.add(f"E{GOAL_FIRST}:E{GOAL_LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="B8A5E0"))
gl.cell(row=GOAL_LAST + 2, column=2, value="Total per month").font = f(10, True)
c = gl.cell(row=GOAL_LAST + 2, column=8, value=f"=SUM(H{GOAL_FIRST}:H{GOAL_LAST})")
c.font = f(12, True, BRAND_DARK)
c.number_format = CUR0
gl.cell(row=GOAL_LAST + 3, column=2, value="Compare this with your 'Savings' budget on Settings.").font = f(9, color=MUTED, italic=True)

# ---------------------------------------------------------------- Debt Payoff
dp = wb.create_sheet("Debt Payoff")
base(dp, {"A": 3, "B": 26, "C": 13, "D": 10, "E": 14, "F": 13, "G": 14, "H": 15, "I": 16}, BAD)
title(dp, "Debt Payoff", "How long each balance takes at the payment you enter, and what it costs.", 8)
dp["B4"] = "Extra monthly payment"
dp["B4"].font = f(10, True)
dp["C4"] = 0 if BLANK else 150
body(dp["C4"], CUR0, is_input=True)
dp["D4"] = "goes to the debt with the highest APR (avalanche method)."
dp["D4"].font = f(9, color=MUTED, italic=True)
header(dp, DEBT_FIRST - 1, ["Debt", "Balance", "APR", "Monthly payment", "Pays extra?", "Months to $0", "Payoff in", "Total interest"])
sample_debts = [] if BLANK else [
    ("Credit card", 4200, 0.2299, 150), ("Car loan", 11800, 0.069, 310), ("Student loan", 18500, 0.048, 210),
]
for i in range(DEBT_FIRST, DEBT_LAST + 1):
    d = sample_debts[i - DEBT_FIRST] if i - DEBT_FIRST < len(sample_debts) else (None, None, None, None)
    for col, val, fmt in ((2, d[0], None), (3, d[1], CUR0), (4, d[2], "0.00%"), (5, d[3], CUR0)):
        body(dp.cell(row=i, column=col, value=val), fmt, is_input=True)
    rng = f"$D${DEBT_FIRST}:$D${DEBT_LAST}"
    # Tie-break on row so exactly one debt takes the extra payment.
    dp.cell(row=i, column=6, value=(
        f'=IF(OR(B{i}="",N(C{i})<=0),"",IF(AND(D{i}=MAX({rng}),COUNTIF($D${DEBT_FIRST}:D{i},D{i})=1),"Yes",""))'))
    pay = f'(E{i}+IF(F{i}="Yes",$C$4,0))'
    rate = f"(D{i}/12)"
    # NPER errors when the payment never covers interest; say so instead of showing #NUM!.
    dp.cell(row=i, column=7, value=(
        f'=IF(OR(B{i}="",N(C{i})<=0),"",IF({pay}<=C{i}*{rate},"Never",'
        f'IF(N(D{i})=0,ROUNDUP(C{i}/{pay},0),ROUNDUP(NPER({rate},-{pay},C{i}),0))))'))
    dp.cell(row=i, column=8, value=f'=IF(ISNUMBER(G{i}),TEXT(DATE(YEAR(TODAY()),MONTH(TODAY())+G{i},1),"mmm yyyy"),"")')
    dp.cell(row=i, column=9, value=(
        f'=IF(ISNUMBER(G{i}),MAX(0,IF(N(D{i})=0,0,{pay}*NPER({rate},-{pay},C{i})-C{i})),"")'))
    for col, fmt, align in ((6, None, "center"), (7, "0", "center"), (8, None, "center"), (9, CUR0, None)):
        body(dp.cell(row=i, column=col), fmt, align=align)
dp.conditional_formatting.add(f"G{DEBT_FIRST}:G{DEBT_LAST}", CellIsRule(operator="equal", formula=['"Never"'], font=Font(name=FONT, bold=True, color=BAD)))
tr = DEBT_LAST + 2
dp.cell(row=tr, column=2, value="Totals").font = f(10, True)
for col, formula, fmt in ((3, f"=SUM(C{DEBT_FIRST}:C{DEBT_LAST})", CUR0), (5, f"=SUM(E{DEBT_FIRST}:E{DEBT_LAST})+C4", CUR0),
                          (7, f'=IF(COUNT(G{DEBT_FIRST}:G{DEBT_LAST})=0,"",MAX(G{DEBT_FIRST}:G{DEBT_LAST}))', "0"),
                          (9, f"=SUM(I{DEBT_FIRST}:I{DEBT_LAST})", CUR0)):
    c = dp.cell(row=tr, column=col, value=formula)
    c.font = f(11, True, BRAND_DARK)
    c.number_format = fmt
dp.cell(row=tr + 1, column=2, value=(
    "Estimates assume a fixed APR and payment. The extra payment is applied to one debt only; "
    "once it's gone, roll its payment into the next highest APR.")).font = f(9, color=MUTED, italic=True)

pie = DoughnutChart()
pie.title = "Where the debt is"
pie.add_data(Reference(dp, min_col=3, min_row=DEBT_FIRST - 1, max_row=DEBT_FIRST + max(len(sample_debts), 3) - 1), titles_from_data=True)
pie.set_categories(Reference(dp, min_col=2, min_row=DEBT_FIRST, max_row=DEBT_FIRST + max(len(sample_debts), 3) - 1))
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
for attr in ("showVal", "showCatName", "showSerName", "showLegendKey"):
    setattr(pie.dataLabels, attr, False)
pie.height, pie.width = 8, 12
dp.add_chart(pie, f"B{tr + 3}")

# Print setup: each tab fits the page width.
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for row in sheet.iter_rows():
        for cell in row:
            if cell.font and cell.font.name != FONT:
                cell.font = Font(name=FONT, size=cell.font.size, bold=cell.font.bold,
                                 italic=cell.font.italic, color=cell.font.color)

wb.active = 1 if not BLANK else 0
# openpyxl writes no cached values; make Excel compute everything on open.
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT)
